from __future__ import annotations
import csv
from pathlib import Path
from typing import Any

def sniff_csv(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()[:65536]
    text = ""
    encoding = "utf-8-sig"
    for candidate in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = raw.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text, delimiters=",;\t|")
        sep = dialect.delimiter
    except csv.Error:
        sep = ";"
    return encoding, sep

def read_csv_rows(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    encoding, sep = sniff_csv(path)
    with path.open("r", encoding=encoding, newline="") as f:
        reader = csv.DictReader(f, delimiter=sep)
        headers = reader.fieldnames or []
        return headers, list(reader)

def _load_openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "Le support Excel nécessite 'openpyxl'.\n"
            "Dans un environnement virtuel : python -m pip install openpyxl\n"
            "Les fichiers CSV ne nécessitent aucune dépendance externe."
        ) from exc

def read_excel_rows(path: Path, options: dict | None = None) -> tuple[list[str], list[dict[str, Any]]]:
    openpyxl = _load_openpyxl()
    wb = openpyxl.load_workbook(path, data_only=False, read_only=False)
    ws = wb.active
    options = options or {}
    header_row = int(options.get("header_row", 1))
    required = options.get("header_columns", [])
    if required:
        matches = [i for i, values in enumerate(ws.iter_rows(values_only=True), 1)
                   if set(required).issubset({str(v).strip() for v in values if v is not None})]
        if len(matches) != 1:
            raise ValueError("En-tête Excel absent ou ambigu : " + ", ".join(required))
        header_row = matches[0]

    headers = ["" if c.value is None else str(c.value) for c in ws[header_row]]
    rows: list[dict[str, Any]] = []

    for row in ws.iter_rows(min_row=header_row+1, values_only=False):
        record: dict[str, Any] = {}
        empty = True
        for idx, cell in enumerate(row):
            if idx >= len(headers):
                break
            value = cell.value
            if value not in (None, ""):
                empty = False
            record[headers[idx]] = value
        if not empty:
            rows.append(record)

    wb.close()
    return headers, rows

def read_source(path: Path, options: dict | None = None) -> tuple[list[str], list[dict[str, Any]]]:
    options = options or {}
    if options.get('adapter'):
        from .profiles import load_adapter
        headers, records = load_adapter(options['adapter']).read_source(path, options)
        rows = list(records)
        if (not isinstance(headers, list) or not all(isinstance(h, str) for h in headers)
                or not all(isinstance(row, dict) for row in rows)):
            raise ValueError('Le lecteur doit retourner une liste de colonnes et des lignes sous forme de dictionnaires.')
        return headers, rows
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv_rows(path)
    if suffix in {".xlsx", ".xlsm"}:
        return read_excel_rows(path, options)
    raise ValueError(f"Format non supporté : {suffix}. Formats : .csv, .xlsx, .xlsm")

def write_csv(path: Path, columns: list[str], rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=columns,
            delimiter=",",
            quoting=csv.QUOTE_MINIMAL,
            extrasaction="ignore",
        )
        writer.writeheader()
        for row in rows:
            writer.writerow({col: row.get(col, "") for col in columns})
